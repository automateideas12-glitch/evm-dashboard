"""
Build Word (.docx) versions of the Markdown reports with compact page settings.
    python reports/build_word_versions.py
Requires pandoc and python-docx. Edit the .md files, then rerun.
"""
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
FILES = sys.argv[1:] or ["user_manual.md", "validation_report.md", "reflection_report_draft.md",
                         "viva_question_bank.md", "presentation_outline.md"]

for name in FILES:
    md = HERE / name
    if not md.exists():
        continue
    out = md.with_suffix(".docx")
    subprocess.run(["pandoc", str(md), "-o", str(out), f"--resource-path={HERE}"], check=True)
    doc = Document(out)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(1.8)
    for style in doc.styles:
        try:
            if style.type == 1:  # paragraph styles
                style.font.name = "Arial"
                if style.name in ("Normal", "Body Text", "First Paragraph", "Compact"):
                    style.font.size = Pt(10.5)
                    style.paragraph_format.space_after = Pt(4)
        except Exception:
            pass
    for t in doc.tables:
        t.style = doc.styles["Table"] if "Table" in [s.name for s in doc.styles] else t.style
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)
    doc.save(out)
    print("built", out.name)
