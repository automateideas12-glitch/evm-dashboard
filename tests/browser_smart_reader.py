"""
Real-browser test of the smart reader upload flow (not collected by pytest).

    streamlit run app.py            # in one terminal
    python tests/browser_smart_reader.py

Uses tests/fixtures/site_status_sheet.xlsx (the team's own site sheet: status date above the table,
header on row 4, fractions, totals row, no actual-cost column).
Path A: type the missing actual cost in the app.  Path B: download the file with the column added,
fill it, upload again.  Both must open the dashboard with CPI 0.954 and SPI 0.504
(EV 10,63,500; AC 11,15,000; PV 21,11,538.46).
"""
import io
import os
import tempfile
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright

URL = os.environ.get("APP_URL", "http://localhost:8501")
ROOT = Path(__file__).resolve().parent.parent
SITE = ROOT / "tests" / "fixtures" / "site_status_sheet.xlsx"
SHOTS = ROOT / "reports" / "screenshots"
TMP = Path(tempfile.gettempdir())
AC = [80000, 125000, 260000, 280000, 240000, 130000, 0, 0, 0, 0, 0, 0]
log, failures = [], []


def note(msg, ok=None):
    line = msg if ok is None else f"{'PASS' if ok else 'FAIL'}: {msg}"
    if ok is False:
        failures.append(msg)
    log.append(line)
    print(line, flush=True)


def idle(page, t=1.5):
    page.wait_for_timeout(400)
    try:
        page.wait_for_selector('[data-testid="stStatusWidget"]', state="detached", timeout=20000)
    except Exception:
        pass
    page.wait_for_timeout(int(t * 1000))


def main_text(page):
    return page.locator("[data-testid='stMain']").first.inner_text()


def metric(page, label):
    t = page.locator('[data-testid="stMetric"]').filter(has_text=label).first.inner_text()
    return [x for x in t.splitlines() if x.strip()][1]


def upload(page, path):
    page.locator('input[type="file"]').first.set_input_files(str(path))
    page.get_by_text("What the smart reader understood").wait_for(timeout=30000)
    idle(page, 2)


def open_dashboard(page, name):
    page.get_by_label("Project name *").first.fill(name)
    page.keyboard.press("Tab")
    idle(page, 2)
    page.get_by_role("button", name="Open dashboard").click()
    page.wait_for_selector(".sb-title", timeout=60000)
    idle(page, 3)


with sync_playwright() as p:
    b = p.chromium.launch(executable_path=os.environ.get("CHROME_PATH") or None)
    page = b.new_page(viewport={"width": 1440, "height": 2600}, accept_downloads=True)

    # ---------------- path A: enter actual cost in the app
    page.goto(URL)
    page.wait_for_selector(".hero-title", timeout=60000)
    idle(page)
    upload(page, SITE)
    body = main_text(page)
    note("data type detected as work-package", "Work-Package-Level Data" in body)
    note("header found on row 4", "header on row 4" in body)
    note("status date read from above the table", "28 Feb 2025" in body)
    note("totals row removed", "Removed the summary row 'Project totals / indices'" in body)
    note("fractions converted", "holds fractions" in body)
    note("missing actual cost reported", "does not contain" in body and "Actual cost to date (cumulative)" in body)
    page.screenshot(path=SHOTS / "08_smart_reader.png", full_page=True)

    grid = page.locator('[data-testid="stDataFrame"]').filter(has_text="Actual cost to date (cumulative)").first
    grid.scroll_into_view_if_needed()
    box = grid.bounding_box()
    x = box["x"] + box["width"] * 0.8          # third column (the grid is drawn on a canvas)
    row_h = 35
    for i, v in enumerate(AC):
        y = box["y"] + row_h * (i + 1) + row_h / 2
        idle(page, 0.5)                      # never type while the app is still updating
        page.mouse.dblclick(x, y)            # open the cell editor
        page.wait_for_timeout(500)
        page.keyboard.press("Control+A")
        page.keyboard.type(str(v), delay=40)
        page.keyboard.press("Enter")
        idle(page, 1.2)
    page.mouse.click(box["x"] + 5, box["y"] - 30)
    idle(page, 2)
    page.get_by_label("Project name *").first.fill("Site sheet (typed AC)")
    page.keyboard.press("Tab")
    idle(page, 3)
    body = main_text(page)
    note("validation passes after typing actual cost", "Validation passed" in body)
    if "Validation passed" in body:
        open_dashboard(page, "Site sheet (typed AC)")
        cpi, spi = metric(page, "CPI (cost efficiency)"), metric(page, "SPI (schedule efficiency)")
        note(f"path A dashboard CPI {cpi} (expected 0.954)", cpi == "0.954")
        note(f"path A dashboard SPI {spi} (expected 0.504)", spi == "0.504")
        page.screenshot(path=SHOTS / "09_smart_reader_dashboard.png", full_page=True)
    else:
        page.screenshot(path=TMP / "smart_reader_typing_failed.png", full_page=True)

    # ---------------- path B: download, fill, upload again
    page.goto(URL)
    page.wait_for_selector(".hero-title", timeout=60000)
    idle(page)
    upload(page, SITE)
    with page.expect_download() as dl:
        page.get_by_role("button", name="Download my file with the missing column(s) added (Excel)").click()
    filled = TMP / "site_status_sheet_to_fill.xlsx"
    dl.value.save_as(filled)
    wb = openpyxl.load_workbook(filled)
    ws = wb["Data"]
    head = [c.value for c in ws[1]]
    col = head.index("Actual cost to date (cumulative)") + 1
    for i, v in enumerate(AC, start=2):
        ws.cell(row=i, column=col, value=v)
    wb.save(filled)
    page.goto(URL)
    page.wait_for_selector(".hero-title", timeout=60000)
    idle(page)
    upload(page, filled)
    page.get_by_label("Project name *").first.fill("Site sheet (filled file)")
    page.keyboard.press("Tab")
    idle(page, 3)
    body = main_text(page)
    note("re-uploaded file: nothing missing", "does not contain" not in body)
    note("re-uploaded file: validation passes", "Validation passed" in body)
    open_dashboard(page, "Site sheet (filled file)")
    cpi, spi = metric(page, "CPI (cost efficiency)"), metric(page, "SPI (schedule efficiency)")
    note(f"path B dashboard CPI {cpi} (expected 0.954)", cpi == "0.954")
    note(f"path B dashboard SPI {spi} (expected 0.504)", spi == "0.504")
    b.close()

note(f"RESULT: {len(failures)} failure(s)")
(ROOT / "reports" / "browser_smart_reader_log.txt").write_text("\n".join(log))
