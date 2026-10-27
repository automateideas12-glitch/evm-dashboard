"""
Smart reader — understands real-world spreadsheets and CSV files before validation.

Rule-based and offline (no AI service). It reads both the HEADERS and the VALUES of a file:

  * finds the header row even when titles, notes or blank rows sit above the table
  * reads label/value cells above the table (e.g. "Status date: 28-Feb-2025") as metadata
  * matches columns to the dashboard's fields by meaning (header words + what the values look like),
    with a confidence level and a reason for every match; unknown columns are never guessed
  * detects the data type (progress / pre-calculated / work-package) from the columns found
  * removes summary rows ("Total", "Grand total", "Project totals / indices")
  * converts % stored as fractions (0.75 = 75 %) and amounts written as lakh / crore
  * lists anything still missing (e.g. actual cost) so the user can supply it

Every change it makes is listed in `Understanding.fixes` and shown to the user; nothing is
changed silently. The cleaned table then goes through the normal validation (prepare()).
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field, replace
from typing import Dict, List, Optional, Tuple

import pandas as pd

from .data_processing import FIELDS, MODE_FIELDS, MODES, SYNONYMS, mode_fields, normalise_header
from .validation import is_missing, to_date, to_number

ENTERED_SUFFIX = " (entered in app)"
ENTRY_LABEL = {"ac_cumulative": "Actual cost to date (cumulative)"}
MIN_SCORE = 0.5
SCAN_ROWS = 30
_SUMMARY = re.compile(r"\b(total|totals|sub ?total|grand total|indices|summary)\b", re.IGNORECASE)
_STATUS_LABEL = re.compile(r"^\s*(status|data|report(ing)?|as ?of|cut ?off|valuation)\s*(date)?\s*:?\s*(.*)$",
                           re.IGNORECASE)
_PROJECT_LABEL = re.compile(r"^\s*project(\s*(name|title))?\s*:\s*(.*)$", re.IGNORECASE)
_AMOUNT_WORD = re.compile(r"^\s*([₹$€£]|rs\.?|inr)?\s*([-\d.,]+)\s*(lakhs?|lacs?|crores?|cr|l)\s*$", re.IGNORECASE)
_WORD_MULT = {"lakh": 1e5, "lakhs": 1e5, "lac": 1e5, "lacs": 1e5, "l": 1e5,
              "crore": 1e7, "crores": 1e7, "cr": 1e7}


# ================================================================ result
@dataclass
class Understanding:
    table: pd.DataFrame                       # cleaned table (original headers + any added columns)
    mode: str
    mode_reason: str
    mapping: Dict[str, Optional[str]]         # field -> column, for every field of the mode
    matches: List[dict]                       # field, column, confidence, reason
    header_row: int                           # spreadsheet row of the header (1-based)
    row_numbers: List[int]                    # spreadsheet row of every table row
    sheet: Optional[str] = None
    metadata: Dict[str, object] = field(default_factory=dict)
    fixes: List[str] = field(default_factory=list)
    hints: List[str] = field(default_factory=list)
    missing: List[str] = field(default_factory=list)
    unused_columns: List[str] = field(default_factory=list)
    checks: List[Tuple[str, str]] = field(default_factory=list)   # ("success"|"warning", message)


# ================================================================ header vocabulary
def _norm(text) -> str:
    s = str(text).lower().replace("%", " pct ").replace("&", " and ")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _has(s: str, pattern: str) -> bool:
    return re.search(pattern, s) is not None


_PCT = r"\b(pct|percent|percentage|progress|complete|completion)\b"
_PLAN = r"\b(planned|plan|baseline|scheduled|schedule|target|budgeted|expected)\b"
_ACTUAL = r"\b(actual|physical|earned|achieved|completed|done|executed)\b"
_MONEYISH = r"\b(value|cost|amount|budget|bac|spent|expenditure)\b"
_TIMEISH = r"\b(date|start|finish|duration|days|month|week)\b"


def header_score(text, fld: str) -> Tuple[float, str]:
    """How well a column header matches a field (0-1) and why. Header words only."""
    if text is None or is_missing(text):
        return 0.0, ""
    canon = normalise_header(text)
    if canon == fld or canon in SYNONYMS.get(fld, []):
        return 1.0, "standard column name"
    s = _norm(text)
    if fld == "reporting_date":
        if _has(s, r"\b(start|finish|baseline|early|late|due|completion|forecast|target|planned|actual)\b"):
            return 0.0, ""
        if _has(s, r"\b(status|data|report|reporting|as of|asof|cut ?off|valuation|progress)\b.*\bdate\b"):
            return 0.95, "header says it is the status/report date"
        if _has(s, r"^(date|month|period|week|period end|period ending|month ending|reporting period)$"):
            return 0.85, "header is a date/period name"
        if _has(s, r"\bdate\b"):
            return 0.6, "header contains 'date'"
        return 0.0, ""
    if fld == "planned_progress_pct":
        if _has(s, _PLAN) and _has(s, _PCT) and not _has(s, _MONEYISH + "|" + _TIMEISH):
            return 0.95, "planned/baseline + % / progress"
        return 0.0, ""
    if fld == "actual_progress_pct":
        if _has(s, _ACTUAL) and _has(s, _PCT) and not _has(s, _MONEYISH + "|" + _TIMEISH + "|" + _PLAN):
            return 0.95, "actual/physical + % / progress"
        return 0.0, ""
    if fld == "pv_cumulative":
        return (0.95, "planned value / BCWS") if _has(s, r"\b(pv|bcws|planned value)\b") and not _has(s, r"\bpct\b") \
            else (0.0, "")
    if fld == "ev_cumulative":
        return (0.95, "earned value / BCWP") if _has(s, r"\b(ev|bcwp|earned value)\b") and not _has(s, r"\bpct\b") \
            else (0.0, "")
    if fld in ("ac_cumulative", "ac_period"):
        if _has(s, r"\b(budget|budgeted|planned|plan|estimate|estimated|pct|forecast|value)\b"):
            return 0.0, ""
        costish = _has(s, r"\b(ac|acwp|actual cost|cost|spent|spend|expenditure|incurred)\b")
        periodic = _has(s, r"\b(this month|monthly|month|period|this period|week|weekly|this)\b")
        if not costish:
            return 0.0, ""
        if fld == "ac_period":
            return (0.95, "cost + this period/month") if periodic else (0.0, "")
        if periodic:
            return 0.0, ""
        if _has(s, r"\b(ac|acwp)\b"):
            return 0.95, "AC / ACWP"
        if _has(s, r"\b(actual cost|cost to date|cost till date|cost incurred|incurred|spent|spend|expenditure)\b"):
            return 0.95, "actual cost / cost to date / spent"
        return 0.55, "header mentions cost"
    if fld == "bac":
        if _has(s, r"\b(actual|spent|pct|earned|planned value)\b"):
            return 0.0, ""
        if _has(s, r"\b(bac|budget at completion|total budget|project budget|contract value|project cost|"
                   r"sanctioned cost|approved budget)\b"):
            return 0.95, "budget at completion"
        return (0.8, "header says budget") if _has(s, r"\bbudget\b") else (0.0, "")
    if fld == "activity_budget":
        if _has(s, r"\b(actual|spent|pct|earned|incurred|planned value)\b"):
            return 0.0, ""
        if _has(s, r"\b(activity budget|wp budget|work package budget|task budget|budgeted cost|item budget|"
                   r"boq amount|boq value)\b"):
            return 0.95, "activity / work-package budget"
        if _has(s, r"\b(budget|bac|contract value|amount)\b"):
            return 0.85, "budget / BAC per row"
        return 0.0, ""
    if fld == "activity_id":
        if _has(s, r"\b(pct|date|cost|budget|value|amount|duration)\b"):
            return 0.0, ""
        if _has(s, r"\b(id|code|wbs|ref)\b") or _has(s, r"^(s|sr|sl|item|activity|task) ?no$"):
            return 0.95, "ID / code / WBS"
        return 0.0, ""
    if fld == "activity_name":
        if _has(s, r"\b(id|code|no|date|pct|cost|budget|value|amount|start|finish|duration)\b"):
            return 0.0, ""
        if _has(s, r"\b(activity|task|description|work item|item|particulars|work package|scope|name)\b"):
            return 0.9, "activity / task / description"
        return 0.0, ""
    if fld == "notes":
        return (0.95, "notes / remarks") if _has(s, r"\b(notes?|remarks?|comments?)\b") else (0.0, "")
    if fld == "project_name":
        return (0.95, "project name") if _has(s, r"^project( name| title)?$") else (0.0, "")
    if fld == "currency":
        return (0.95, "currency") if _has(s, r"^(currency|curr)$") else (0.0, "")
    if fld == "independent_etc":
        return (0.9, "estimate to complete") if _has(s, r"\b(etc|estimate to complete)\b") else (0.0, "")
    return 0.0, ""


# ================================================================ value helpers
def _amount(v):
    """Number from a cell, accepting '137 lakh' / '4 crore'. Returns (value, used_word)."""
    if isinstance(v, str):
        m = _AMOUNT_WORD.match(v)
        if m:
            try:
                return float(m.group(2).replace(",", "")) * _WORD_MULT[m.group(3).lower()], True
            except ValueError:
                return None, False
    x, err = to_number(v, allow_percent_sign=True)
    return (None if err else x), False


def _content_ok(values: List, kind: str) -> bool:
    vals = [v for v in values if not is_missing(v)]
    if not vals:
        return True                       # an empty column says nothing against the header
    if kind == "date":
        ok = sum(to_date(v)[0] is not None for v in vals)
    elif kind == "pct":
        nums = [_amount(v)[0] for v in vals]
        ok = sum(x is not None and 0 <= x <= 100 for x in nums)
    elif kind == "money":
        ok = sum(_amount(v)[0] is not None for v in vals)
    else:
        return True
    return ok / len(vals) >= 0.8


# ================================================================ reading the grid
@dataclass
class _Grid:
    rows: List[List[object]]
    pct_format: Dict[Tuple[int, int], bool]
    uncached_formula: Dict[Tuple[int, int], bool]
    name: Optional[str] = None


def _csv_grid(data: bytes) -> _Grid:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = "\n".join(text.splitlines()[:50])
    delim = max([",", ";", "\t", "|"], key=lambda d: sum(line.count(d) for line in sample.splitlines()))
    rows = [[(c if c.strip() != "" else None) for c in r] for r in csv.reader(io.StringIO(text), delimiter=delim)]
    return _Grid(rows, {}, {})


def _xlsx_grids(data: bytes) -> List[_Grid]:
    import openpyxl
    wb_v = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    wb_f = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    grids = []
    for ws in wb_v.worksheets:
        wf = wb_f[ws.title]
        rows, pct, unc = [], {}, {}
        for r_i, row in enumerate(ws.iter_rows()):
            out = []
            for c_i, cell in enumerate(row):
                out.append(cell.value)
                if "%" in (cell.number_format or ""):
                    pct[(r_i, c_i)] = True
                if cell.value is None:
                    f = wf.cell(row=cell.row, column=cell.column).value
                    if isinstance(f, str) and f.startswith("="):
                        unc[(r_i, c_i)] = True
            rows.append(out)
        grids.append(_Grid(rows, pct, unc, ws.title))
    return grids


def _row_header_quality(row: List[object]) -> Tuple[int, float]:
    hits, total = 0, 0.0
    for v in row:
        if isinstance(v, str):
            best = max(header_score(v, f)[0] for f in FIELDS)
            if best >= MIN_SCORE:
                hits += 1
                total += best
    return hits, total


def _find_header(grid: _Grid) -> Tuple[int, Tuple[int, float]]:
    best_i, best_q = None, (0, 0.0)
    for i, row in enumerate(grid.rows[:SCAN_ROWS]):
        q = _row_header_quality(row)
        if q > best_q:
            best_i, best_q = i, q
    if best_i is None or best_q[0] < 2:
        best_i = next((i for i, r in enumerate(grid.rows) if any(not is_missing(v) for v in r)), 0)
    return best_i, best_q


def _col_letter(i: int) -> str:
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def _metadata(grid: _Grid, header_i: int) -> Dict[str, object]:
    meta: Dict[str, object] = {}
    for r_i, row in enumerate(grid.rows[:header_i]):
        for c_i, v in enumerate(row):
            if not isinstance(v, str):
                continue
            def value_after():
                inline = v.split(":", 1)[1].strip() if ":" in v else ""
                if inline:
                    return inline, f"{_col_letter(c_i)}{r_i + 1}"
                for c2 in range(c_i + 1, len(row)):
                    if not is_missing(row[c2]):
                        return row[c2], f"{_col_letter(c2)}{r_i + 1}"
                return None, None
            m = _STATUS_LABEL.match(v)
            if m and "status_date" not in meta and (m.group(3) or ":" in v):
                val, ref = value_after()
                d, err = to_date(val) if val is not None else (None, "none")
                if d is not None:
                    meta["status_date"], meta["status_date_cell"] = d, ref
                    continue
            if _PROJECT_LABEL.match(v) and "project_name" not in meta:
                val, _ = value_after()
                if val is not None:
                    meta["project_name"] = str(val).strip()
    return meta


# ================================================================ mapping
def _assign(table: pd.DataFrame, fields: List[str]) -> Tuple[Dict[str, Optional[str]], List[dict]]:
    cands = []
    for col in table.columns:
        vals = list(table[col])
        for f in fields:
            sc, why = header_score(col, f)
            if sc < MIN_SCORE:
                continue
            kind = FIELDS[f][2]
            if not _content_ok(vals, kind):
                continue                    # header looked right but the values do not fit
            cands.append((sc, f, col, why))
    cands.sort(key=lambda t: (-t[0], fields.index(t[1])))
    mapping: Dict[str, Optional[str]] = {f: None for f in fields}
    matches, used = [], set()
    for sc, f, col, why in cands:
        if mapping[f] is None and col not in used:
            mapping[f] = col
            used.add(col)
            conf = "High" if sc >= 0.9 else "Medium" if sc >= 0.7 else "Low"
            kind = FIELDS[f][2]
            reason = why + {"date": "; values are dates", "pct": "; values are 0-100 %",
                            "money": "; values are amounts"}.get(kind, "")
            matches.append({"field": f, "column": col, "confidence": conf, "reason": reason})
    return mapping, matches


def _detect_mode(table: pd.DataFrame) -> Tuple[str, str]:
    m, _ = _assign(table, mode_fields("work_package"))
    if m["activity_id"] and m["activity_budget"] and (m["planned_progress_pct"] or m["actual_progress_pct"]):
        return "work_package", (f"Columns '{m['activity_id']}' and '{m['activity_budget']}' show one row per "
                                "activity / work package with its own budget.")
    m, _ = _assign(table, mode_fields("precalc"))
    if m["pv_cumulative"] and m["ev_cumulative"]:
        return "precalc", f"The file already has PV ('{m['pv_cumulative']}') and EV ('{m['ev_cumulative']}')."
    m, _ = _assign(table, mode_fields("progress"))
    if m["planned_progress_pct"] or m["actual_progress_pct"]:
        return "progress", "One row per reporting date with planned and actual progress %."
    return "progress", "Could not tell the data type from the columns; assumed one row per reporting date."


def _missing(mode: str, mapping: Dict[str, Optional[str]], table: pd.DataFrame) -> List[str]:
    spec = MODE_FIELDS[mode]
    out = [f for f in spec["required"] if not mapping.get(f)]
    if spec["ac_choice"] and not mapping.get("ac_cumulative") and not mapping.get("ac_period"):
        out.append("ac_cumulative")
    # a column mapped but completely empty is still missing data
    out += [f for f, c in mapping.items() if c and f in spec["required"] + ["ac_cumulative"]
            and f not in out and all(is_missing(v) for v in table[c])]
    return out


# ================================================================ cross-check of supplied PV / EV
def _cross_check(table: pd.DataFrame, mapping: Dict[str, Optional[str]], mode: str, unused: List[str],
                 hints: List[str]) -> Tuple[List[Tuple[str, str]], List[str]]:
    """
    Outside pre-calculated mode the dashboard derives PV = budget x planned % and EV = budget x actual %.
    If the file also has PV / EV columns, compare them row by row instead of ignoring them.
    Returns (checks, columns that matched). Columns that differ are reported, and stay a validation warning.
    """
    if mode == "precalc":
        return [], []
    budget_col = mapping.get("activity_budget") if mode == "work_package" else mapping.get("bac")
    if mode == "work_package":
        key_col, unit = mapping.get("activity_id"), "activities"
    else:
        key_col, unit = mapping.get("reporting_date"), "reporting dates"
    checks, matched, ok_labels = [], [], []
    for fld, pct_field, label, word in (("pv_cumulative", "planned_progress_pct", "PV", "planned"),
                                        ("ev_cumulative", "actual_progress_pct", "EV", "actual")):
        col = next((c for c in unused if header_score(c, fld)[0] >= MIN_SCORE), None)
        if col is None:
            continue
        pct_col = mapping.get(pct_field)
        if not budget_col or not pct_col:
            hints.append(f"Your '{col}' column is not used: the dashboard calculates {label} itself from the budget "
                         f"and {word} progress %.")
            continue
        diffs, compared = [], 0
        for i in range(len(table)):
            supplied = _amount(table[col].iloc[i])[0]
            budget = _amount(table[budget_col].iloc[i])[0]
            pct = _amount(table[pct_col].iloc[i])[0]
            if supplied is None or budget is None or pct is None:
                continue
            compared += 1
            calc = budget * pct / 100.0
            if abs(supplied - calc) > max(0.5, 1e-6 * abs(calc)):
                key = table[key_col].iloc[i] if key_col else f"row {i + 1}"
                if isinstance(key, pd.Timestamp):
                    key = f"{key:%d %b %Y}"
                diffs.append(f"{key} (file {supplied:,.0f}; calculated {calc:,.0f})")
        if not compared:
            continue
        if diffs:
            more = f" and {len(diffs) - 5} more" if len(diffs) > 5 else ""
            checks.append(("warning", f"Your {label} column '{col}' differs from budget × {word} % for "
                                      f"{len(diffs)} of {compared} {unit}: {'; '.join(diffs[:5])}{more}. The "
                                      f"dashboard uses budget × {word} %; check the {word} % or the {label} formula "
                                      "in your file."))
        else:
            matched.append(col)
            ok_labels.append(label)
    if ok_labels:
        cols = " and ".join(ok_labels)
        checks.insert(0, ("success", f"Your {cols} column{'s' if len(ok_labels) > 1 else ''} match"
                                     f"{'' if len(ok_labels) > 1 else 'es'} the dashboard's own calculation "
                                     f"(budget × progress %) for all {len(table)} {unit}."))
    return checks, matched


# ================================================================ main entry
def understand(file_bytes: bytes, filename: str, mode: Optional[str] = None,
               sheet: Optional[str] = None) -> Understanding:
    name = filename.lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        grids = [_csv_grid(file_bytes)]
    elif name.endswith(".xlsx"):
        grids = _xlsx_grids(file_bytes)
        if sheet is not None:
            grids = [g for g in grids if g.name == sheet] or grids
    else:
        raise ValueError("Unsupported file type. Upload a .csv or .xlsx file.")

    # choose the sheet whose header row looks most like EVM data
    scored = [(g, *_find_header(g)) for g in grids]
    grid, h_i, _ = max(scored, key=lambda t: t[2])

    fixes: List[str] = []
    hints: List[str] = []
    header = grid.rows[h_i] if h_i < len(grid.rows) else []
    width = max((len(r) for r in grid.rows), default=0)
    body = [list(r) + [None] * (width - len(r)) for r in grid.rows[h_i + 1:]]
    header = list(header) + [None] * (width - len(header))

    # columns: keep headed columns and unheaded columns that contain data
    cols, names, seen = [], [], {}
    for c_i, h in enumerate(header):
        has_data = any(not is_missing(r[c_i]) for r in body)
        if is_missing(h) and not has_data:
            continue
        nm = str(h).strip() if not is_missing(h) else f"Column {_col_letter(c_i)}"
        if nm in seen:
            seen[nm] += 1
            nm = f"{nm} ({seen[nm]})"
        else:
            seen[nm] = 1
        cols.append(c_i)
        names.append(nm)

    rows, row_numbers, summary, blank = [], [], [], 0
    for r_off, r in enumerate(body):
        vals = [r[c] for c in cols]
        if all(is_missing(v) for v in vals):
            blank += 1
            continue
        first = next((v for v in vals if not is_missing(v)), None)
        if isinstance(first, str) and _SUMMARY.search(first):
            summary.append((first.strip(), h_i + 2 + r_off))
            continue
        rows.append(vals)
        row_numbers.append(h_i + 2 + r_off)
    table = pd.DataFrame(rows, columns=names)
    if h_i > 0:
        fixes.append(f"Found the column headers on row {h_i + 1} (the rows above are titles/notes, not data).")
    for text, rn in summary:
        fixes.append(f"Removed the summary row '{text}' (row {rn}); the dashboard calculates totals itself.")

    meta = _metadata(grid, h_i)
    detected_mode, reason = _detect_mode(table)
    use_mode = mode or detected_mode
    if mode and mode != detected_mode:
        reason = f"Data type chosen by you. (Detected: {MODES[detected_mode]}.)"
    fields = mode_fields(use_mode)
    mapping, matches = _assign(table, fields)

    # status date written above the table -> one reporting date for every row
    if not mapping.get("reporting_date") and "status_date" in meta:
        col = f"Status date (from cell {meta['status_date_cell']})"
        table[col] = meta["status_date"]
        mapping["reporting_date"] = col
        matches.insert(0, {"field": "reporting_date", "column": col, "confidence": "High",
                           "reason": "status date written above the table"})
        fixes.append(f"Used the status date {meta['status_date']:%d %b %Y} from cell {meta['status_date_cell']} "
                     "as the reporting date of every row.")

    # value fixes on mapped columns
    for f, col in mapping.items():
        if not col:
            continue
        kind = FIELDS[f][2]
        c_i = cols[names.index(col)] if col in names else None
        if kind == "pct":
            nums = [_amount(v)[0] for v in table[col]]
            present = [x for x in nums if x is not None]
            pct_fmt = c_i is not None and any(grid.pct_format.get((h_i + 1 + i, c_i)) for i in range(len(body)))
            fractional = any(0 < x < 1 for x in present)
            if present and max(present) <= 1 and (pct_fmt or fractional) and not any(
                    isinstance(v, str) and "%" in v for v in table[col]):
                table[col] = [None if x is None else x * 100 for x in nums]
                fixes.append(f"'{col}' holds fractions (0.75 = 75%); converted to percentages.")
        elif kind == "money":
            conv = [_amount(v) for v in table[col]]
            if any(w for _, w in conv):
                table[col] = [x if w else v for (x, w), v in zip(conv, table[col])]
                fixes.append(f"'{col}' had amounts written in lakh/crore; converted to full amounts.")
        if c_i is not None and any(grid.uncached_formula.get((h_i + 1 + i, c_i)) for i in range(len(body))):
            hints.append(f"'{col}' contains formulas whose results were not saved in the file. Open the file in "
                         "Excel, press Save, and upload it again.")

    used = {c for c in mapping.values() if c}
    unused = [c for c in table.columns if c not in used]
    for c in unused:
        s = _norm(c)
        if _has(s, r"^(pct|percent) complete$|^pct$|^progress$"):
            hints.append(f"'{c}' was not used as physical progress: a bare '% Complete' (e.g. from MS Project) is "
                         "usually duration-based. Map it by hand only if it really measures physical work.")
    checks, crosschecked = _cross_check(table, mapping, use_mode, unused, hints)
    table.attrs["crosschecked_columns"] = crosschecked
    if meta.get("project_name"):
        pass
    table.attrs["row_numbers"] = row_numbers
    table.attrs["blank_rows_skipped"] = blank
    u = Understanding(table=table, mode=use_mode, mode_reason=reason, mapping=mapping, matches=matches,
                      header_row=h_i + 1, row_numbers=row_numbers, sheet=grid.name, metadata=meta, fixes=fixes,
                      hints=hints, unused_columns=unused, checks=checks)
    u.missing = _missing(use_mode, mapping, table)
    return u


# ================================================================ supplying missing values
def entry_label(fld: str) -> str:
    return ENTRY_LABEL.get(fld, FIELDS[fld][0])


def _key_columns(u: Understanding) -> List[str]:
    keys = [u.mapping.get("activity_id"), u.mapping.get("activity_name")] if u.mode == "work_package" \
        else [u.mapping.get("reporting_date")]
    return [k for k in keys if k]


def values_entry_table(u: Understanding, fld: str) -> pd.DataFrame:
    """A small table (one row per data row) for the user to type a missing field's values."""
    keys = _key_columns(u)
    out = u.table[keys].copy().reset_index(drop=True)
    out[entry_label(fld)] = [None] * len(out)
    return out


def with_entered_values(u: Understanding, fld: str, entry: pd.DataFrame) -> Understanding:
    lab = entry_label(fld)
    vals = list(entry[lab]) if lab in entry else []
    if len(vals) != len(u.table) or all(is_missing(v) for v in vals):
        return u
    table = u.table.copy()
    table.attrs = dict(u.table.attrs)
    col = lab + ENTERED_SUFFIX
    table[col] = vals
    mapping = dict(u.mapping)
    mapping[fld] = col
    matches = [m for m in u.matches if m["field"] != fld] + [
        {"field": fld, "column": col, "confidence": "High", "reason": "entered in the app"}]
    new = replace(u, table=table, mapping=mapping, matches=matches,
                  fixes=u.fixes + [f"{lab} entered in the app for {sum(not is_missing(v) for v in vals)} row(s)."])
    new.missing = _missing(new.mode, mapping, table)
    return new


def workbook_with_missing_columns(u: Understanding) -> bytes:
    """The cleaned table as .xlsx with empty columns for every missing field, ready to fill in and re-upload."""
    t = u.table.copy()
    for f in u.missing:
        t[entry_label(f)] = None
    bio = io.BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as xw:
        t.to_excel(xw, sheet_name="Data", index=False)
    return bio.getvalue()
