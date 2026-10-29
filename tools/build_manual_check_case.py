"""
Build the MANUAL-CHECK CASE: a real project small enough to verify every dashboard number by hand.

Source: IJzertoren Memorial Square (DSLIB C2015-33, Belgium, 2013-14), converted with
tools/convert_dslib.py into data/real/C2015-33_ijzertoren_memorial_square_work_packages.csv
(12 activities x 5 month-end progress reports, ProTrack's recorded PV/EV/AC).

What this script does (no data is invented):
  1. groups the 12 activities into 4 work packages by trade (table WORK_PACKAGES below):
       budget  = sum of activity budgets
       planned % = sum of activity PV / WP budget x 100      (PV is preserved exactly)
       actual %  = sum of activity EV / WP budget x 100      (EV is preserved exactly)
       AC        = sum of activity cumulative AC
  2. rounds budgets and AC to cents and percentages to 2 decimals so the case can be worked with a
     calculator. Effect: project PV/EV differ from ProTrack by at most EUR 3; CPI/SPI agree to 3 decimals.
  3. writes
       data/real/C2015-33_ijzertoren_manual_check_work_packages.csv   (upload in the dashboard)
       data/manual_check/IJzertoren_manual_check_workbook.xlsx        (hand calculation, Excel formulas)

The workbook contains formulas only (no pasted results). Open it in Excel and save once so the cached
values are stored (tests read the cached values); the copy in the repository was recalculated with
LibreOffice.

    python tools/build_manual_check_case.py
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "real" / "C2015-33_ijzertoren_memorial_square_work_packages.csv"
OUT_CSV = ROOT / "data" / "real" / "C2015-33_ijzertoren_manual_check_work_packages.csv"
OUT_XLSX = ROOT / "data" / "manual_check" / "IJzertoren_manual_check_workbook.xlsx"
PROJECT = "IJzertoren Memorial Square (manual check case)"

# activity -> work package (DSLIB activity names are Dutch; translations in brackets)
WORK_PACKAGES = {
    "A01": ("WP1", "Demolition and sewerage"),                    # opbraak & riolering
    "A02": ("WP2", "Earthworks and foundations"),                 # grondwerken & onderfundering (zone 1)
    "A03": ("WP2", "Earthworks and foundations"),                 # fundering (zone 1)
    "A08": ("WP2", "Earthworks and foundations"),                 # grondwerken & onderfundering (zone 2)
    "A09": ("WP2", "Earthworks and foundations"),                 # fundering (zone 2)
    "A04": ("WP3", "Kerbs, paving and concrete"),                 # boordstenen (zone 1)
    "A05": ("WP3", "Kerbs, paving and concrete"),                 # verharding
    "A06": ("WP3", "Kerbs, paving and concrete"),                 # bekisting + beton gieten (zone 1)
    "A10": ("WP3", "Kerbs, paving and concrete"),                 # boordstenen (zone 2)
    "A11": ("WP3", "Kerbs, paving and concrete"),                 # bekisting + beton gieten (zone 2)
    "A07": ("WP4", "Finishing: street furniture and signage"),    # afwerking (zone 1)
    "A12": ("WP4", "Finishing: street furniture and signage"),    # afwerking (zone 2)
}
BASELINE_START, BASELINE_FINISH = "14-Oct-2013", "20-Dec-2013"


def build_csv() -> pd.DataFrame:
    a = pd.read_csv(SRC)
    missing = set(a["activity_id"]) - set(WORK_PACKAGES)
    if missing:
        raise SystemExit(f"Activities without a work package: {missing}")
    a["wp_id"] = a["activity_id"].map(lambda i: WORK_PACKAGES[i][0])
    a["wp_name"] = a["activity_id"].map(lambda i: WORK_PACKAGES[i][1])
    a["pv"] = a["activity_budget"] * a["planned_progress_pct"] / 100
    a["ev"] = a["activity_budget"] * a["actual_progress_pct"] / 100
    g = (a.groupby(["reporting_date", "wp_id", "wp_name"], sort=True)
          .agg(budget=("activity_budget", "sum"), pv=("pv", "sum"), ev=("ev", "sum"), ac=("ac_cumulative", "sum"))
          .reset_index())
    out = pd.DataFrame({
        "project_name": PROJECT,
        "reporting_date": g["reporting_date"],
        "activity_id": g["wp_id"],
        "activity_name": g["wp_name"],
        "activity_budget": g["budget"].round(2),
        "planned_progress_pct": (g["pv"] / g["budget"] * 100).round(2),
        "actual_progress_pct": (g["ev"] / g["budget"] * 100).round(2),
        "ac_cumulative": g["ac"].round(2),
        "notes": "",
    })
    OUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    out.to_csv(OUT_CSV, index=False)
    return out


# ---------------------------------------------------------------- workbook
ARIAL = "Arial"
HEAD = Font(name=ARIAL, bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F2A33")
INPUT = Font(name=ARIAL, size=10, color="0000FF")      # blue = typed input
CALC = Font(name=ARIAL, size=10)                        # black = formula
KEY = PatternFill("solid", fgColor="FFF4D6")            # highlight: the two benchmark periods
EUR = '€#,##0.00;-€#,##0.00;"-"'
IDX = "0.0000"
PCT = '0.00"%"'


def header(ws, cols, widths):
    for j, (h, w) in enumerate(zip(cols, widths), 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill = HEAD, FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def put(ws, r, c, v, font=CALC, fmt=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    if fmt:
        cell.number_format = fmt
    return cell


def build_workbook(df: pd.DataFrame):
    wb = Workbook()
    dates = sorted(df["reporting_date"].unique())
    period_of = {d: i + 1 for i, d in enumerate(dates)}
    n = len(df)
    last = n + 1

    # ---- Read me
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 120
    lines = [
        ("Manual check case: IJzertoren Memorial Square (real project)", True),
        ("Real project from the DSLIB database (Ghent University): Batselier, J. & Vanhoucke, M. (2015), "
         "International Journal of Project Management 33(3), 697-710. Project C2015-33, Belgium, 2013-14, EUR.", False),
        (f"Baseline {BASELINE_START} to {BASELINE_FINISH}; 5 month-end progress reports; finished on 28-Feb-2014, "
         "about 2 months late and 4.8 % over budget.", False),
        ("The 12 real activities are grouped into 4 work packages. Budgets, PV and EV are sums of the activity values "
         "recorded in ProTrack; % values are rounded to 2 decimals so everything can be checked with a calculator.", False),
        ("", False),
        ("Sheets", True),
        ("Inputs - the 20 input rows (5 dates x 4 work packages), identical to the upload file "
         "C2015-33_ijzertoren_manual_check_work_packages.csv. Blue = typed input.", False),
        ("Work package calculation - PV, EV, AC, CV, SV, CPI, SPI and status for every work package and date.", False),
        ("Project calculation - every dashboard metric for every date: CPI, SPI, EAC (A, B, C), ETC, VAC, TCPI, "
         "health. Rows 3 and 4 (30-Nov and 31-Dec) are the two benchmark periods worked by hand in the validation report.", False),
        ("Forecast check - each forecast compared with the real final cost (known because the project is complete).", False),
        ("Feature tour - which dashboard feature to show at which date, using this data.", False),
        ("Settings - health thresholds (same defaults as the dashboard sidebar) and an example for method D.", False),
        ("", False),
        ("How to check the app: upload the CSV (or press 'Manual check case' on the dashboard home screen), pick a "
         "reporting date in the sidebar, and compare each figure with the same row here. All cells in black are "
         "formulas - click any cell to see how it is calculated.", False),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(name=ARIAL, bold=bold, size=13 if i == 1 else 10)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    # ---- Inputs
    ws = wb.create_sheet("Inputs")
    header(ws, ["Period", "Reporting date", "WP ID", "Work package", "Budget", "Planned progress %",
                "Actual progress %", "AC cumulative"], [8, 14, 8, 36, 14, 12, 12, 15])
    for i, r in enumerate(df.itertuples(index=False), 2):
        put(ws, i, 1, period_of[r.reporting_date], INPUT)
        put(ws, i, 2, pd.Timestamp(r.reporting_date).to_pydatetime(), INPUT, "dd-mmm-yyyy")
        put(ws, i, 3, r.activity_id, INPUT)
        put(ws, i, 4, r.activity_name, INPUT)
        put(ws, i, 5, float(r.activity_budget), INPUT, EUR)
        put(ws, i, 6, float(r.planned_progress_pct), INPUT, PCT)
        put(ws, i, 7, float(r.actual_progress_pct), INPUT, PCT)
        put(ws, i, 8, float(r.ac_cumulative), INPUT, EUR)

    # ---- Settings
    st = wb.create_sheet("Settings")
    st.column_dimensions["A"].width = 62
    st.column_dimensions["B"].width = 16
    st.column_dimensions["C"].width = 70
    settings = [
        ("Setting", "Value", "Note"),
        ("CPI/SPI warning below", 1.00, "Dashboard default (sidebar: Health screening thresholds)"),
        ("CPI/SPI critical below", 0.90, "Dashboard default"),
        ("Forecast overrun warning above (%)", 0.0, "Dashboard default"),
        ("Forecast overrun critical above (%)", 10.0, "Dashboard default"),
        ("Actual final cost (AC at completion, 28-Feb-2014)", "='Project calculation'!F6",
         "Known only because the project is finished; used in 'Forecast check'"),
        ("Method D: independent ETC you want to test (enter a number)", None,
         "Example: type a bottom-up estimate of the remaining cost at 31-Dec-2013"),
        ("Method D: EAC = AC (31-Dec-2013) + independent ETC", '=IF(ISNUMBER(B7),\'Project calculation\'!F4+B7,"Enter an ETC in B7")',
         "Compare with the dashboard: Forecasts tab > Independent estimate, date 31 Dec 2013"),
    ]
    for i, row in enumerate(settings, 1):
        for j, v in enumerate(row, 1):
            c = st.cell(row=i, column=j, value=v)
            c.font = HEAD if i == 1 else (INPUT if j == 2 and i in (2, 3, 4, 5, 7) else CALC)
            if i == 1:
                c.fill = FILL
            if j == 2 and i in (6, 8):
                c.number_format = EUR

    # ---- Work package calculation (one row per Inputs row)
    ws = wb.create_sheet("Work package calculation")
    cols = ["Period", "Status date", "WP ID", "Work package", "Budget", "Planned %", "Actual %", "PV", "EV", "AC",
            "CV", "SV", "CPI", "SPI", "Status"]
    header(ws, cols, [8, 13, 7, 34, 13, 10, 10, 13, 13, 13, 13, 13, 9, 9, 16])
    for r in range(2, last + 1):
        f = {
            1: (f"=Inputs!A{r}", None), 2: (f"=Inputs!B{r}", "dd-mmm-yyyy"), 3: (f"=Inputs!C{r}", None),
            4: (f"=Inputs!D{r}", None), 5: (f"=Inputs!E{r}", EUR), 6: (f"=Inputs!F{r}", PCT), 7: (f"=Inputs!G{r}", PCT),
            8: (f"=E{r}*F{r}/100", EUR), 9: (f"=E{r}*G{r}/100", EUR), 10: (f"=Inputs!H{r}", EUR),
            11: (f"=I{r}-J{r}", EUR), 12: (f"=I{r}-H{r}", EUR),
            13: (f'=IF(J{r}>0,I{r}/J{r},"Undefined")', IDX), 14: (f'=IF(H{r}>0,I{r}/H{r},"Undefined")', IDX),
            15: (f'=IF(AND(H{r}<=0.005,I{r}<=0.005,J{r}<=0.005),"Not started",'
                 f'IF(OR(ISTEXT(M{r}),ISTEXT(N{r})),"Insufficient data",'
                 f'IF(MIN(M{r},N{r})<Settings!$B$3,"Critical",'
                 f'IF(MIN(M{r},N{r})<Settings!$B$2,"Warning","Favourable"))))', None),
        }
        for c, (fx, fmt) in f.items():
            put(ws, r, c, fx, CALC, fmt)

    # ---- Project calculation (one row per reporting date)
    ws = wb.create_sheet("Project calculation")
    cols = ["Period", "Status date", "BAC", "PV", "EV", "AC", "CV", "SV", "CPI", "SPI", "% planned", "% complete",
            "% spent", "EAC A", "EAC B", "EAC C", "ETC A", "ETC B", "ETC C", "VAC A", "VAC B", "VAC C",
            "Overrun % A", "TCPI to BAC", "TCPI to EAC A", "TCPI to EAC B", "TCPI to EAC C", "CPI status",
            "SPI status", "Overrun status (A)", "Overall health (method A)"]
    header(ws, cols, [7, 12, 13, 13, 13, 13, 12, 12, 8, 8, 9, 9, 9, 13, 13, 13, 13, 13, 13, 12, 12, 12, 10, 9, 9,
                      9, 9, 12, 12, 12, 14])
    W = "'Work package calculation'"
    rng = lambda col: f"{W}!${col}$2:${col}${last}"  # noqa: E731
    for i in range(1, len(dates) + 1):
        r = i + 1
        put(ws, r, 1, i, INPUT)
        f = {
            2: (f"=INDEX({rng('B')},MATCH(A{r},{rng('A')},0))", "dd-mmm-yyyy"),
            3: (f"=SUMIFS({rng('E')},{rng('A')},A{r})", EUR),
            4: (f"=SUMIFS({rng('H')},{rng('A')},A{r})", EUR),
            5: (f"=SUMIFS({rng('I')},{rng('A')},A{r})", EUR),
            6: (f"=SUMIFS({rng('J')},{rng('A')},A{r})", EUR),
            7: (f"=E{r}-F{r}", EUR), 8: (f"=E{r}-D{r}", EUR),
            9: (f'=IF(F{r}>0,E{r}/F{r},"Undefined")', IDX), 10: (f'=IF(D{r}>0,E{r}/D{r},"Undefined")', IDX),
            11: (f"=D{r}/C{r}*100", "0.00"), 12: (f"=E{r}/C{r}*100", "0.00"), 13: (f"=F{r}/C{r}*100", "0.00"),
            14: (f'=IF(ISNUMBER(I{r}),IF(I{r}>0,C{r}/I{r},"Undefined"),"Undefined")', EUR),
            15: (f'=IF(AND(ISNUMBER(I{r}),ISNUMBER(J{r})),IF(I{r}*J{r}>0,F{r}+(C{r}-E{r})/(I{r}*J{r}),"Undefined"),'
                 f'"Undefined")', EUR),
            16: (f"=F{r}+C{r}-E{r}", EUR),
        }
        for k, (eac_col, letter) in enumerate((("N", "N"), ("O", "O"), ("P", "P"))):
            f[17 + k] = (f"=IF(ISNUMBER({letter}{r}),{letter}{r}-F{r},{letter}{r})", EUR)        # ETC
            f[20 + k] = (f"=IF(ISNUMBER({letter}{r}),C{r}-{letter}{r},{letter}{r})", EUR)        # VAC
        f[23] = (f"=IF(ISNUMBER(N{r}),(N{r}-C{r})/C{r}*100,N{r})", "0.00")
        f[24] = (f'=IF(C{r}-E{r}<1E-9,"N/A",IF(C{r}-F{r}<1E-9,"Undefined",(C{r}-E{r})/(C{r}-F{r})))', IDX)
        for k, letter in enumerate(("N", "O", "P")):
            f[25 + k] = (f'=IF(C{r}-E{r}<1E-9,"N/A",IF(NOT(ISNUMBER({letter}{r})),"Undefined",'
                         f'IF({letter}{r}-F{r}<1E-9,"Undefined",(C{r}-E{r})/({letter}{r}-F{r}))))', IDX)
        idx_status = lambda c: (f'=IF(ISNUMBER({c}{r}),IF({c}{r}<Settings!$B$3,"Critical",'  # noqa: E731
                                f'IF({c}{r}<Settings!$B$2,"Warning","Favourable")),"Insufficient data")')
        f[28] = (idx_status("I"), None)
        f[29] = (idx_status("J"), None)
        f[30] = (f'=IF(ISNUMBER(W{r}),IF(W{r}>Settings!$B$5,"Critical",IF(W{r}>Settings!$B$4,"Warning",'
                 f'"Favourable")),"Insufficient data")', None)
        f[31] = (f'=IF(OR(AB{r}="Critical",AC{r}="Critical",AD{r}="Critical"),"Critical",'
                 f'IF(OR(AB{r}="Insufficient data",AC{r}="Insufficient data",AD{r}="Insufficient data"),'
                 f'"Insufficient data",IF(OR(AB{r}="Warning",AC{r}="Warning",AD{r}="Warning"),"Warning",'
                 f'"Favourable")))', None)
        for c, (fx, fmt) in f.items():
            cell = put(ws, r, c, fx, CALC, fmt)
            if i in (2, 3):
                cell.fill = KEY
        if i in (2, 3):
            ws.cell(row=r, column=1).fill = KEY

    # ---- Forecast check
    ws = wb.create_sheet("Forecast check")
    header(ws, ["Period", "Status date", "% complete", "EAC A", "EAC B", "EAC C", "Actual final cost",
                "Error A", "Error B", "Error C"], [7, 13, 11, 14, 14, 14, 15, 10, 10, 10])
    P = "'Project calculation'"
    for i in range(1, len(dates) + 1):
        r = i + 1
        f = {1: (f"={P}!A{r}", None), 2: (f"={P}!B{r}", "dd-mmm-yyyy"), 3: (f"={P}!L{r}", "0.00"),
             4: (f"={P}!N{r}", EUR), 5: (f"={P}!O{r}", EUR), 6: (f"={P}!P{r}", EUR), 7: ("=Settings!$B$6", EUR)}
        for k, col in enumerate(("D", "E", "F")):
            f[8 + k] = (f"=IF(ISNUMBER({col}{r}),{col}{r}/G{r}-1,{col}{r})", "+0.0%;-0.0%;0.0%")
        for c, (fx, fmt) in f.items():
            put(ws, r, c, fx, CALC, fmt)
    note = ws.cell(row=len(dates) + 3, column=1,
                   value="Error = forecast / actual final cost - 1. Negative = the method under-forecast the final cost. "
                         "Read row 4 (31-Dec-2013, about 52 % complete) as 'what would we have predicted half way?'")
    note.font = Font(name=ARIAL, size=9, italic=True)

    # ---- Feature tour
    ws = wb.create_sheet("Feature tour")
    header(ws, ["Feature", "Reporting date to select", "What to show / explain"], [32, 22, 110])
    tour = [
        ("Home screen and upload", "-", "Upload the CSV, or press 'Manual check case' on the home screen. The smart reader "
         "shows the data type it detected and the column matches."),
        ("Signboard and KPI cards", "31 Oct 2013", "CPI 1.135 (cost favourable) but SPI 0.544 (behind schedule) -> Critical."),
        ("Not started / early start", "31 Oct 2013 and 30 Nov 2013",
         "31 Oct: WP3 and WP4 'Not started'. 30 Nov: WP4 has progress before any work was planned (PV = 0), "
         "so its SPI is undefined -> 'Insufficient data' (never shown as 0)."),
        ("Top problem activities", "30 Nov 2013", "WP2 Earthworks and foundations: AC far above EV (CPI about 0.37)."),
        ("CPI and SPI trend (combined chart)", "28 Feb 2014", "CPI falls from 1.14 to 0.71 and recovers to 0.95; "
         "SPI recovers to 1.00 at completion although the project was 2 months late (SPI limitation)."),
        ("S-curve", "any", "PV reaches BAC on 20 Dec 2013 (baseline finish); EV and AC continue to 28 Feb 2014."),
        ("Forecasts A, B, C", "31 Dec 2013", "Half way (52 % complete): compare EAC A, B, C with the real final cost "
         "EUR 224,789.67 (sheet 'Forecast check')."),
        ("Method D (independent ETC)", "31 Dec 2013", "Forecasts tab: enter an ETC; the EAC equals AC + ETC (Settings B8)."),
        ("TCPI", "31 Dec 2013 and 28 Feb 2014", "31 Dec: efficiency needed to finish on budget. 28 Feb: EV = BAC, "
         "so TCPI does not apply (shown as N/A)."),
        ("Health thresholds", "any", "Change 'CPI/SPI critical below' in the sidebar and in Settings B3: both update."),
        ("Commentary and exports", "any", "Commentary tab text; Data tab downloads (CSV, Excel workbook)."),
    ]
    for i, row in enumerate(tour, 2):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = CALC
            c.alignment = Alignment(wrap_text=True, vertical="top")

    wb.move_sheet("Settings", offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("Settings"))
    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)


if __name__ == "__main__":
    data = build_csv()
    build_workbook(data)
    print(f"wrote {OUT_CSV.relative_to(ROOT)} ({len(data)} rows) and {OUT_XLSX.relative_to(ROOT)}")
