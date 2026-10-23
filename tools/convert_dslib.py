"""
Convert a REAL project from the DSLIB empirical project database (Ghent University,
OR&S group) into the EVM Dashboard's CSV input formats.

Data source (cite when used):
    Batselier, J. & Vanhoucke, M. (2015). Construction and evaluation framework for a
    real-life project database. International Journal of Project Management, 33(3),
    697-710. https://doi.org/10.1016/j.ijproman.2014.09.004
    Download: https://github.com/mariovanhoucke/DSLIB-Dynamic-Scheduling-Empirical-Project-Library/releases/latest
    (unzip DSLIB<version>.zip; project workbooks are in the "Excel" folder)

Each DSLIB workbook has one sheet per tracking period ("Project Control - TP1", "TP2", ...)
listing every activity with: baseline Total Cost (budget), cumulative Actual Cost,
Percentage Completed, Earned Value and Planned Value, as recorded in ProTrack.

Output (written next to --out prefix):
    <prefix>_work_packages.csv   -> dashboard mode "Work-Package-Level Data"
    <prefix>_evm.csv             -> dashboard mode "Pre-calculated EVM Data"
                                    (project totals exactly as reported by ProTrack)
    <prefix>_protrack_reference.csv -> ProTrack's own CPI/SPI per period, for cross-checking

Conversion rules (no data is invented):
    activity_budget       = baseline "Total Cost"
    actual_progress_pct   = activity EV / activity budget x 100  (ProTrack's earned value)
                            NOT the "Percentage Completed" column: in ProTrack that is a DURATION
                            percentage (elapsed / total days), and ProTrack earns an activity's fixed
                            cost at its start. Using it would understate EV (by up to EUR 32,405 in
                            C2014-05) - the duration-vs-physical-progress trap.
    planned_progress_pct  = activity PV / activity budget x 100  (ProTrack's time-phased PV)
    ac_cumulative         = "Actual Cost" (cumulative to the status date)
    Activities with a zero budget have no PV/EV; their planned % is taken as their actual %
    (it does not affect PV or EV) and they are reported in the console output.

Usage:
    python tools/convert_dslib.py "DSLIB 3.4/Excel/C2014-05 Apartment Building (1).xlsx" --out data/real_C2014-05
"""
import argparse
import csv
import re
import sys
from pathlib import Path

import openpyxl

COLS = {"id": "ID", "name": "Name", "budget": "Total Cost", "ac": "Actual Cost",
        "pct": "Percentage Completed", "ev": "Earned Value (EV)", "pv": "Planned Value (PV)"}


def tp_number(name):
    m = re.search(r"TP(\d+)$", name)
    return int(m.group(1)) if m else None


def read_tp(ws):
    rows = list(ws.iter_rows(values_only=True))
    status = rows[0][2]
    header = next(i for i, r in enumerate(rows) if r and r[0] == "ID")
    h = list(rows[header])
    idx = {}
    for key, label in COLS.items():
        idx[key] = h.index(label)  # first occurrence (the money columns, not the % block)
    acts = []
    total = None
    for r in rows[header + 1:]:
        if r[idx["id"]] is None:
            continue
        rec = {k: r[i] for k, i in idx.items()}
        if rec["id"] == 0:
            total = rec          # project summary row
        else:
            acts.append(rec)
    return status, acts, total


def is_summary(rec, all_ids, wbs_parent):
    return rec["id"] in wbs_parent


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--out", required=True, help="output path prefix")
    ap.add_argument("--currency", default="EUR")
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.xlsx, read_only=True, data_only=True)
    base = list(wb["Baseline Schedule"].iter_rows(values_only=True))
    project_name = base[2][1]
    # WBS summary rows (parents) would double count: identify activities that have children
    wbs = {r[0]: str(r[2]) for r in base[2:] if r and r[0] is not None}
    parents = {i for i, w in wbs.items() if any(o != w and o.startswith(w + ".") for o in wbs.values())}

    tps = sorted((s for s in wb.sheetnames if tp_number(s)), key=tp_number)
    if not tps:
        sys.exit("This project has no tracking periods - choose one with tracking data.")
    wp_rows, evm_rows, ref_rows, zero_budget = [], [], [], set()
    for s in tps:
        status, acts, total = read_tp(wb[s])
        date = status.date().isoformat()
        for r in acts:
            if r["id"] in parents:
                continue
            budget = float(r["budget"] or 0)
            if budget > 0:
                pct_actual = float(r["ev"] or 0) / budget * 100
                pct_plan = float(r["pv"] or 0) / budget * 100
            else:
                pct_actual = float(r["pct"] or 0) * 100
                pct_plan = pct_actual
                zero_budget.add(f"{r['id']} {r['name']}")
            wp_rows.append([project_name, date, f"A{r['id']:02d}", r["name"], round(budget, 2),
                            round(pct_plan, 8), round(pct_actual, 8), round(float(r["ac"] or 0), 2), ""])
        evm_rows.append([project_name, date, round(float(total["budget"]), 2), round(float(total["pv"]), 2),
                         round(float(total["ev"]), 2), round(float(total["ac"]), 2), f"DSLIB {s}"])
    ov = list(wb["Tracking Overview"].iter_rows(values_only=True))
    hdr = list(ov[1])
    for r in ov[2:]:
        if r[0] is None:
            continue
        d = dict(zip(hdr, r))
        ref_rows.append([r[0], d["Status date"].date().isoformat(), d["Planned Value (PV)"], d["Earned Value (EV)"],
                         d["Actual Cost (AC)"], d["Schedule Performance Index (SPI)"],
                         d["Cost Performance Index (CPI)"]])

    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    with open(f"{out}_work_packages.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["project_name", "reporting_date", "activity_id", "activity_name", "activity_budget",
                    "planned_progress_pct", "actual_progress_pct", "ac_cumulative", "notes"])
        w.writerows(wp_rows)
    with open(f"{out}_evm.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["project_name", "reporting_date", "bac", "pv_cumulative", "ev_cumulative", "ac_cumulative",
                    "notes"])
        w.writerows(evm_rows)
    with open(f"{out}_protrack_reference.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["period_name", "status_date", "PV", "EV", "AC", "SPI_protrack", "CPI_protrack"])
        w.writerows(ref_rows)
    print(f"{project_name}: {len(tps)} tracking periods, {len(wp_rows) // len(tps)} activities per period, "
          f"{len(parents)} WBS summary rows skipped, currency {a.currency}")
    if zero_budget:
        print("Zero-budget activities (planned % set equal to actual %; no effect on PV/EV):", sorted(zero_budget))


if __name__ == "__main__":
    main()
